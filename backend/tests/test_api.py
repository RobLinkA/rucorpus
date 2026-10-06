"""API tests using invented records only; no historical or production dataset."""
import io
import json
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from corpus.models import AuditLog, Corpus, Segment, SegmentAnnotation

pytestmark = [pytest.mark.django_db, pytest.mark.usefixtures('dataset')]


def put(c, url, data):
    return c.put(url, json.dumps(data), content_type='application/json')


def post(c, url, data):
    return c.post(url, json.dumps(data), content_type='application/json')


def test_search_annotation_logic(anon):
    or_ = anon.get('/api/search', {'values': [1, 3]}).json()['total']
    and_ = anon.get('/api/search', {'values': [1, 3], 'logic': 'and'}).json()['total']
    one = anon.get('/api/search', {'values': [1]}).json()['total']
    assert or_ >= one >= and_


def test_search_requires_condition(anon):
    assert anon.get('/api/search', {'q': ''}).status_code == 400


def test_hidden_corpus_invisible_to_public(anon, admin):
    corpus = admin.get('/api/manage/corpora').json()[1]
    put(admin, f"/api/manage/corpora/{corpus['id']}", {**corpus, 'status': 'hidden'})
    try:
        names = [c['name'] for c in anon.get('/api/meta').json()['corpora']]
        assert corpus['name'] not in names
        assert anon.get('/api/search', {'q': 'сказал', 'corpora': [corpus['id']]}).json()['total'] == 0
    finally:
        put(admin, f"/api/manage/corpora/{corpus['id']}", {**corpus, 'status': 'published'})


def test_detail_navigation(anon):
    g = anon.get('/api/read/1', {'size': 10}).json()['groups'][1]
    d = anon.get(f"/api/groups/{g['id']}").json()
    assert d['prev'] and d['next'] and any(c['is_current'] for c in d['context'])


def test_permissions(anon, reader, editor):
    seg = Segment.objects.filter(corpus_id=1).first()
    assert anon.get('/api/manage/dashboard').status_code == 401
    assert reader.get('/api/manage/dashboard').status_code == 403
    assert put(reader, f'/api/segments/{seg.id}/annotations', {'values': []}).status_code == 403
    assert editor.get('/api/manage/users').status_code == 403
    assert editor.get('/api/manage/dashboard').status_code == 200


def test_csrf_required(reader):
    from django.test import Client
    c = Client(enforce_csrf_checks=True)
    c.cookies = reader.cookies
    assert post(c, '/api/me/favorites', {'group_id': 1}).status_code == 403


def test_login_requires_csrf():
    from django.test import Client
    c = Client(enforce_csrf_checks=True)
    c.get('/api/auth/session')
    credentials = {'username': 'reader', 'password': 'test-only-password-42'}
    assert post(c, '/api/auth/login', credentials).status_code == 403
    c.defaults['HTTP_X_CSRFTOKEN'] = c.cookies['csrftoken'].value
    assert post(c, '/api/auth/login', credentials).status_code == 200


def test_https_proxy_keeps_secure_requests_without_redirect(reader, settings):
    settings.SECURE_SSL_REDIRECT = True
    settings.SECURE_PROXY_SSL_HEADER = ('HTTP_X_FORWARDED_PROTO', 'https')
    settings.SESSION_COOKIE_SECURE = settings.CSRF_COOKIE_SECURE = True
    reader.defaults.update(HTTP_HOST='corpus.example.org', HTTP_ORIGIN='https://corpus.example.org',
                           HTTP_X_FORWARDED_PROTO='https')
    response = reader.get('/api/auth/session')
    assert response.status_code == 200
    assert response.cookies['csrftoken']['secure']
    assert post(reader, '/api/me/favorites', {'group_id': 1}).status_code == 200


@pytest.mark.parametrize('host', ['localhost:5173', '127.0.0.1:5173'])
def test_favorites_and_history(reader, host):
    # The dev proxy must preserve Host so it matches the browser's Origin.
    reader.defaults.update(HTTP_HOST=host, HTTP_ORIGIN=f'http://{host}')
    r = reader.get('/api/search', {'q': 'reader'}).json()
    gid = r['results'][0]['id']
    assert post(reader, '/api/me/favorites', {'group_id': gid}).status_code == 200
    assert put(reader, f'/api/me/favorites/by-group/{gid}', {'note': '好例子'}).status_code == 200
    favs = reader.get('/api/me/favorites').json()
    assert favs['total'] == 1 and favs['results'][0]['note'] == '好例子'
    assert reader.get('/api/me/favorites/export').status_code == 200
    assert reader.get('/api/me/history').json()['total'] >= 1
    assert reader.get(f'/api/groups/{gid}').json()['favorite'] is True
    assert reader.delete(f'/api/me/favorites/by-group/{gid}').status_code == 200
    assert reader.get('/api/me/favorites').json()['total'] == 0
    assert reader.get(f'/api/groups/{gid}').json()['favorite'] is False


@pytest.mark.parametrize(('host', 'origin'), [
    ('127.0.0.1:8000', 'http://localhost:5173'),  # the old proxy rewrote Host
    ('localhost:5173', 'https://untrusted.example'),
])
def test_favorites_reject_cross_origin_even_with_valid_token(reader, host, origin):
    reader.defaults.update(HTTP_HOST=host, HTTP_ORIGIN=origin)
    assert post(reader, '/api/me/favorites', {'group_id': 1}).status_code == 403
    assert reader.get('/api/me/favorites').json()['total'] == 0


def test_editor_annotates_with_audit(editor):
    seg = Segment.objects.filter(corpus_id=1).exclude(annotations=None).first()
    before = list(seg.annotations.values_list('id', flat=True))
    r = put(editor, f'/api/segments/{seg.id}/annotations', {'values': before[:-1]})
    assert r.status_code == 200
    assert AuditLog.objects.filter(action='annotate', target=f'segment:{seg.id}').exists()
    put(editor, f'/api/segments/{seg.id}/annotations', {'values': before})
    # Values owned by another corpus must be rejected.
    from corpus.models import AnnotationGroup, AnnotationValue
    group = AnnotationGroup.objects.create(corpus_id=2, name='Foreign test group')
    value = AnnotationValue.objects.create(group=group, label='Foreign test label')
    assert put(editor, f'/api/segments/{seg.id}/annotations', {'values': [value.id]}).status_code == 400


def test_split_unsplit_segment(editor):
    seg = Segment.objects.first()
    seg.unsplit = True
    seg.save()
    parts = [{'source': 'Первая часть.', 'target': '第一部分。'}, {'source': 'Вторая часть.', 'target': '第二部分。'}]
    r = post(editor, f'/api/manage/segments/{seg.id}/split', {'parts': parts})
    assert r.status_code == 200 and len(r.json()['ids']) == 2
    assert not Segment.objects.get(id=seg.id).unsplit


def test_annotation_import_roundtrip(admin):
    from openpyxl import load_workbook
    content = admin.get('/api/manage/transfer/export/1').content
    wb = load_workbook(io.BytesIO(content))
    ws = wb['句对']
    header = [c.value for c in ws[1]]
    col = header.index('翻译技巧') + 1
    sid = ws.cell(2, 1).value
    ws.cell(2, col).value = '减译'
    ws.cell(3, col).value = '没有这个值'
    buf = io.BytesIO()
    wb.save(buf)
    job = admin.post('/api/manage/transfer/upload', {
        'kind': 'annotations', 'corpus_id': 1, 'file': SimpleUploadedFile('a.xlsx', buf.getvalue())}).json()
    assert job['preview']['summary']['changed_segments'] == 1
    assert job['preview']['error_count'] == 1
    done = admin.post(f"/api/manage/transfer/jobs/{job['id']}/commit").json()
    assert done['status'] == 'done' and done['result']['skipped_rows'] == 1
    labels = set(Segment.objects.get(id=sid).annotations.filter(group__name='翻译技巧').values_list('label', flat=True))
    assert labels == {'减译'}


def test_unchanged_workbook_preserves_ai_rows(admin):
    before = list(SegmentAnnotation.objects.filter(origin='ai').values_list(
        'id', 'method', 'evidence', 'confidence', 'replaced_id'))
    content = admin.get('/api/manage/transfer/export/1').content
    response = admin.post('/api/manage/transfer/upload', {
        'kind': 'annotations', 'corpus_id': 1,
        'file': SimpleUploadedFile('synthetic.xlsx', content)})
    assert response.status_code == 200
    job = response.json()
    assert job['preview']['error_count'] == 0
    assert job['preview']['summary']['changed_segments'] == 0
    assert admin.post(f"/api/manage/transfer/jobs/{job['id']}/commit").status_code == 200
    assert before == list(SegmentAnnotation.objects.filter(origin='ai').values_list(
        'id', 'method', 'evidence', 'confidence', 'replaced_id'))


def test_import_new_corpus_preserves_language_metadata_and_rebuilds_index(admin, anon):
    from openpyxl import load_workbook
    content = admin.get('/api/manage/transfer/export/2').content
    workbook = load_workbook(io.BytesIO(content))
    for row in workbook['语料库'].iter_rows(min_row=2):
        if row[0].value == '名称':
            row[1].value = 'Synthetic import copy'
    buffer = io.BytesIO()
    workbook.save(buffer)
    before = Corpus.objects.count()
    job = admin.post('/api/manage/transfer/upload', {
        'kind': 'corpus', 'file': SimpleUploadedFile('synthetic.xlsx', buffer.getvalue())}).json()
    assert job['preview']['error_count'] == 0
    assert Corpus.objects.count() == before  # Preview must not write texts.
    done = admin.post(f"/api/manage/transfer/jobs/{job['id']}/commit").json()
    corpus = Corpus.objects.get(id=done['result']['corpus_id'])
    assert (corpus.source_lang, corpus.target_lang) == ('en', 'fr')
    assert corpus.status == 'hidden' and corpus.versions.filter(lang='fr').count() == 2
    assert Segment.objects.filter(corpus=corpus).count() == 4
    assert anon.get('/api/search', {'q': 'école', 'corpora': [corpus.id]}).json()['total'] == 0
    corpus.status = 'published'
    corpus.save()
    assert anon.get('/api/search', {'q': 'école', 'corpora': [corpus.id]}).json()['total'] == 1
