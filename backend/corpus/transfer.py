"""Excel / CSV exchange format.

The 句对 sheet is self-describing, so the same columns work as a CSV:

    句对ID | 作品 | 段落 | 句组 | 版本 | 句序 | 原文 | 译文 | <标注组 A> | <标注组 B> | ...

* 段落 / 句组 / 句序 are 1-based. Rows of different 版本 sharing (作品, 段落, 句组) are aligned.
* Annotation cells hold value labels separated by '；' (or ';' / ',' / '，').
* Annotation import matches rows by 句对ID; corpus import ignores it and creates everything.

Optional sheets (xlsx only): 语料库 (key/value), 版本, 作品 — add metadata for corpus import.
"""
import csv
import io
import re
from collections import defaultdict
from dataclasses import dataclass, field

from django.db import transaction
from django.db.models import Max

from . import fts
from .models import (AlignGroup, AnnotationGroup, AnnotationValue, Corpus, Paragraph, Segment, SegmentAnnotation,
                     Version, Work)

FIXED = ['句对ID', '作品', '段落', '句组', '版本', '句序', '原文', '译文']
AI_MARK = '✦'  # prefix for AI annotations in exported cells (ignored on import)
SPLIT_RE = re.compile(r'[；;，,、\n]+')
KIND_LABEL = {'source': '原文', 'translation': '译文'}


# ------------------------------------------------------------------ export
def corpus_workbook(corpus: Corpus, with_data=True, work_ids=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    head_font, head_fill = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='4F46E5')

    def sheet(title, rows, widths=None):
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
        for i, w in enumerate(widths or [], start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        ws.freeze_panes = 'A2'
        return ws

    wb.remove(wb.active)
    sheet('说明', [['说明'], *[[line] for line in __doc__.strip().splitlines()]], [110])
    sheet('语料库', [['字段', '值'], ['名称', corpus.name], ['说明', corpus.description],
                    ['原文语言', corpus.source_lang], ['译文语言', corpus.target_lang]], [14, 80])
    sheet('版本', [['名称', '类型', '语言', '作者/译者', '出处', '排序']] +
          [[v.label, KIND_LABEL[v.kind], v.lang, v.person, v.bibliography, v.sort_order]
           for v in corpus.versions.all()], [24, 8, 8, 20, 60, 6])
    works = corpus.works.all() if not work_ids else corpus.works.filter(id__in=work_ids)
    sheet('作品', [['原文标题', '译文标题', '作者', '排序']] +
          [[w.title_source, w.title_target, w.author, w.sort_order] for w in works], [30, 30, 16, 6])
    groups = list(corpus.annotation_groups.prefetch_related('values'))
    sheet('标注体系', [['标注组', '标注项', '控件']] +
          [[g.name, v.label, g.widget] for g in groups for v in g.values.all()], [16, 30, 12])
    rows = [FIXED + [g.name for g in groups]]
    if with_data:
        rows += segment_rows(corpus, groups, works)
    ws = sheet('句对', rows, [10, 22, 6, 6, 18, 6, 70, 60] + [22] * len(groups))
    for col in ('G', 'H'):
        for c in ws[col][1:]:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def segment_rows(corpus, groups, works):
    value_group = {v.id: (g.id, v.label) for g in groups for v in g.values.all()}
    gidx = {g.id: i for i, g in enumerate(groups)}
    anns = defaultdict(list)
    ai = set()
    for sid, vid, origin in SegmentAnnotation.objects.filter(segment__corpus=corpus) \
            .order_by('value__sort_order').values_list('segment_id', 'value_id', 'origin'):
        anns[sid].append(vid)
        if origin == 'ai':
            ai.add((sid, vid))
    title = {w.id: w.title_source for w in works}
    rows = []
    qs = Segment.objects.filter(work_id__in=title).select_related('group__paragraph', 'version') \
        .order_by('work__sort_order', 'group__position', 'version__sort_order', 'seq')
    for s in qs.iterator(chunk_size=2000):
        cells = [[] for _ in groups]
        for vid in anns.get(s.id, []):
            g, label = value_group[vid]
            cells[gidx[g]].append(f'{AI_MARK}{label}' if (s.id, vid) in ai else label)
        rows.append([s.id, title[s.work_id], s.group.paragraph.seq + 1, s.group.seq + 1, s.version.label,
                     s.seq + 1, s.source, s.target] + ['；'.join(c) for c in cells])
    return rows


def csv_bytes(rows):
    buf = io.StringIO()
    buf.write('﻿')
    csv.writer(buf).writerows(rows)
    return buf.getvalue().encode('utf-8')


# ------------------------------------------------------------------ reading files
class FileError(ValueError):
    pass


def read_tables(name, content):
    """Return {sheet_name: [[cell, ...], ...]} for xlsx or csv input."""
    lower = name.lower()
    if lower.endswith('.csv'):
        for enc in ('utf-8-sig', 'gb18030'):
            try:
                text = content.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise FileError('无法识别 CSV 编码，请另存为 UTF-8')
        return {'句对': [row for row in csv.reader(io.StringIO(text))]}
    if lower.endswith('.xlsx'):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            raise FileError(f'无法读取 Excel 文件：{e}')
        return {ws.title: [['' if c is None else c for c in row] for row in ws.iter_rows(values_only=True)]
                for ws in wb.worksheets}
    raise FileError('只支持 .xlsx 或 .csv 文件')


def _s(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _int(v):
    v = _s(v)
    return int(v) if re.fullmatch(r'\d+', v) else None


@dataclass
class Parsed:
    header: list
    rows: list = field(default_factory=list)        # (excel_row_number, dict)
    errors: list = field(default_factory=list)      # (row, message)


def parse_pairs(tables):
    sheet = tables.get('句对')
    if sheet is None:
        raise FileError('缺少“句对”工作表')
    rows = [r for r in sheet if any(_s(c) for c in r)]
    if not rows:
        raise FileError('“句对”工作表为空')
    header = [_s(c) for c in rows[0]]
    missing = [h for h in ('原文',) if h not in header]
    if missing:
        raise FileError(f"表头缺少列：{'、'.join(missing)}（应为：{' | '.join(FIXED)} | 标注组…）")
    p = Parsed(header=header)
    for i, r in enumerate(rows[1:], start=2):
        p.rows.append((i, {h: _s(r[j]) if j < len(r) else '' for j, h in enumerate(header) if h}))
    return p


def _annotation_columns(header, groups):
    names = {g.name: g for g in groups}
    return {h: names[h] for h in header if h in names}, [h for h in header if h and h not in FIXED and h not in names]


def _parse_cell(cell, group, errors, rownum):
    out = set()
    labels = {v.label: v.id for v in group.values.all()}
    for label in SPLIT_RE.split(cell or ''):
        label = label.strip().lstrip(AI_MARK).strip()
        if not label:
            continue
        if label in labels:
            out.add(labels[label])
        else:
            errors.append((rownum, f'“{group.name}”中没有标注项“{label}”'))
    if group.single_choice and len(out) > 1:
        errors.append((rownum, f'“{group.name}”为单选，只能填一个值'))
    return out


# ------------------------------------------------------------------ annotation import
def plan_annotations(corpus: Corpus, tables, update_text=False):
    p = parse_pairs(tables)
    if '句对ID' not in p.header:
        raise FileError('标注导入需要“句对ID”列（请先从本系统导出数据再修改）')
    groups = list(corpus.annotation_groups.prefetch_related('values'))
    cols, unknown_cols = _annotation_columns(p.header, groups)
    if not cols and not update_text:
        raise FileError('没有找到与该语料库标注组同名的列')
    ids = [_int(r.get('句对ID')) for _, r in p.rows]
    segs = {s.id: s for s in Segment.objects.filter(corpus=corpus, id__in=[i for i in ids if i])}
    current = defaultdict(set)
    group_of_value = {v.id: g.id for g in groups for v in g.values.all()}
    for sid, vid in SegmentAnnotation.objects.filter(segment_id__in=segs).values_list('segment_id', 'value_id'):
        current[sid].add(vid)
    changes, adds, removes, texts, unchanged = [], 0, 0, 0, 0
    seen = set()
    for rownum, r in p.rows:
        sid = _int(r.get('句对ID'))
        if sid is None or sid not in segs:
            p.errors.append((rownum, f"句对ID“{r.get('句对ID')}”不存在于该语料库"))
            continue
        if sid in seen:
            p.errors.append((rownum, f'句对ID {sid} 重复'))
            continue
        seen.add(sid)
        before = current[sid]
        touched_groups = {g.id for g in cols.values()}
        after = {v for v in before if group_of_value[v] not in touched_groups}
        n_errors = len(p.errors)
        for h, g in cols.items():
            after |= _parse_cell(r.get(h, ''), g, p.errors, rownum)
        if len(p.errors) > n_errors:  # skip rows with invalid cells entirely
            continue
        text = None
        if update_text:
            src, tgt = r.get('原文', ''), r.get('译文', '')
            if src and (src != segs[sid].source or tgt != segs[sid].target):
                text = (src, tgt)
                texts += 1
        if after != before or text:
            changes.append({'id': sid, 'add': sorted(after - before), 'remove': sorted(before - after),
                            'text': text})
            adds += len(after - before)
            removes += len(before - after)
        else:
            unchanged += 1
    summary = {'rows': len(p.rows), 'changed_segments': len(changes), 'unchanged': unchanged,
               'annotations_added': adds, 'annotations_removed': removes, 'text_updates': texts,
               'columns': list(cols), 'ignored_columns': unknown_cols}
    return summary, changes, p.errors


def apply_annotations(changes):
    with transaction.atomic():
        for c in changes:
            if c['remove']:
                SegmentAnnotation.objects.filter(segment_id=c['id'], value_id__in=c['remove']).delete()
        SegmentAnnotation.objects.bulk_create([SegmentAnnotation(segment_id=c['id'], value_id=v)
                                               for c in changes for v in c['add']], batch_size=2000)
        text_rows = [(c['id'], *c['text']) for c in changes if c['text']]
        for sid, src, tgt in text_rows:
            Segment.objects.filter(id=sid).update(source=src, target=tgt)
        if text_rows:
            fts.index_segments(text_rows)


# ------------------------------------------------------------------ corpus import
def _kv(tables, name):
    out = {}
    for r in tables.get(name, [])[1:]:
        if r and _s(r[0]):
            out[_s(r[0])] = _s(r[1]) if len(r) > 1 else ''
    return out


def _records(tables, name):
    rows = [r for r in tables.get(name, []) if any(_s(c) for c in r)]
    if not rows:
        return []
    header = [_s(c) for c in rows[0]]
    return [{h: _s(r[j]) if j < len(r) else '' for j, h in enumerate(header)} for r in rows[1:]]


def plan_corpus(tables, target: Corpus | None):
    """Validate a corpus workbook. Returns (summary, plan, errors)."""
    p = parse_pairs(tables)
    for h in ('作品', '版本'):
        if h not in p.header:
            raise FileError(f'语料导入需要“{h}”列')
    info = _kv(tables, '语料库')
    if target is None and not info.get('名称'):
        raise FileError('新建语料库需要“语料库”工作表中的“名称”')
    if target is None and Corpus.objects.filter(name=info['名称']).exists():
        raise FileError(f"语料库“{info['名称']}”已存在；如要追加作品，请选择目标语料库")
    version_meta = {r['名称']: r for r in _records(tables, '版本') if r.get('名称')}
    work_meta = {r['原文标题']: r for r in _records(tables, '作品') if r.get('原文标题')}
    scheme = defaultdict(list)
    for r in _records(tables, '标注体系'):
        if r.get('标注组') and r.get('标注项'):
            scheme[r['标注组']].append(r['标注项'])
    existing_groups = {g.name: g for g in target.annotation_groups.prefetch_related('values')} if target else {}
    group_values = {name: [v.label for v in g.values.all()] for name, g in existing_groups.items()}
    for name, labels in scheme.items():
        group_values.setdefault(name, [])
        group_values[name] += [x for x in labels if x not in group_values[name]]
    ann_cols = [h for h in p.header if h and h not in FIXED]
    for h in ann_cols:
        group_values.setdefault(h, [])
    existing_versions = {v.label: v for v in target.versions.all()} if target else {}
    existing_works = set(target.works.values_list('title_source', flat=True)) if target else set()

    units = defaultdict(lambda: defaultdict(list))   # work -> (para, group) -> [(version, seq, src, tgt, anns)]
    versions_seen, new_values = [], defaultdict(set)
    for rownum, r in p.rows:
        work, version, src = r.get('作品', ''), r.get('版本', ''), r.get('原文', '')
        para, grp, seq = _int(r.get('段落')), _int(r.get('句组')), _int(r.get('句序'))
        if not work or not version or not src:
            p.errors.append((rownum, '作品、版本、原文不能为空'))
            continue
        if para is None or para < 1:
            p.errors.append((rownum, '段落必须是正整数'))
            continue
        if work in existing_works:
            p.errors.append((rownum, f'作品“{work}”已存在于目标语料库'))
            continue
        if version not in versions_seen:
            versions_seen.append(version)
        anns = {}
        for h in ann_cols:
            labels = [x.strip().lstrip(AI_MARK).strip() for x in SPLIT_RE.split(r.get(h, '')) if x.strip()]
            for x in labels:
                if x not in group_values[h]:
                    new_values[h].add(x)
            if labels:
                anns[h] = labels
        units[work][(para, grp)].append({'version': version, 'seq': seq, 'source': src,
                                         'target': r.get('译文', ''), 'anns': anns, 'row': rownum})
    # alignment: rows without 句组 are aligned by 句序 within the paragraph
    plan_works = []
    for work, cells in units.items():
        paras = defaultdict(lambda: defaultdict(list))
        for (para, grp), items in cells.items():
            for it in items:
                key = grp if grp is not None else (it['seq'] or 0)
                paras[para][key].append(it)
        plan_works.append({'title': work, 'paragraphs': [
            [sorted(paras[pn][k], key=lambda it: (it['version'], it['seq'] or 0)) for k in sorted(paras[pn])]
            for pn in sorted(paras)]})
    n_segments = sum(len(g) for w in plan_works for para in w['paragraphs'] for g in para)
    summary = {
        'target': target.name if target else info.get('名称'), 'new_corpus': target is None,
        'works': len(plan_works), 'segments': n_segments,
        'groups': sum(len(para) for w in plan_works for para in w['paragraphs']),
        'versions': [{'label': v, 'new': v not in existing_versions} for v in versions_seen],
        'annotation_columns': ann_cols,
        'new_annotation_values': {k: sorted(v) for k, v in new_values.items()},
    }
    versions_all = list(version_meta) + [v for v in versions_seen if v not in version_meta]
    summary['versions'] = [{'label': v, 'new': v not in existing_versions, 'rows': v in versions_seen}
                           for v in versions_all]
    plan = {'info': info, 'version_meta': version_meta, 'work_meta': work_meta, 'versions': versions_all,
            'group_values': {k: v + sorted(new_values.get(k, [])) for k, v in group_values.items()},
            'works': plan_works}
    return summary, plan, p.errors


def apply_corpus(plan, target: Corpus | None):
    info = plan['info']
    with transaction.atomic():
        corpus = target or Corpus.objects.create(
            name=info['名称'], description=info.get('说明', ''), source_lang=info.get('原文语言', 'ru') or 'ru',
            target_lang=info.get('译文语言', 'zh') or 'zh', status=Corpus.Status.HIDDEN,
            sort_order=(Corpus.objects.aggregate(m=Max('sort_order'))['m'] or 0) + 1)
        versions = {v.label: v for v in corpus.versions.all()}
        for i, label in enumerate(plan['versions']):
            if label in versions:
                continue
            m = plan['version_meta'].get(label, {})
            kind = 'source' if m.get('类型') == '原文' else 'translation'
            versions[label] = Version.objects.create(
                corpus=corpus, kind=kind, label=label, lang=m.get('语言') or corpus.target_lang,
                person=m.get('作者/译者', ''), bibliography=m.get('出处', ''),
                sort_order=_int(m.get('排序')) or len(versions) + 1)
        groups = {g.name: g for g in corpus.annotation_groups.all()}
        values = {}
        for gi, (name, labels) in enumerate(plan['group_values'].items()):
            g = groups.get(name) or AnnotationGroup.objects.create(corpus=corpus, name=name, sort_order=gi)
            have = {v.label: v for v in g.values.all()}
            for vi, label in enumerate(labels):
                if label not in have:
                    have[label] = AnnotationValue.objects.create(group=g, label=label, sort_order=vi,
                                                                 defined_in_legacy=False)
                values[(name, label)] = have[label].id
        base_order = (corpus.works.aggregate(m=Max('sort_order'))['m'] or 0) + 1
        new_seg_ids = []
        for wi, w in enumerate(plan['works']):
            m = plan['work_meta'].get(w['title'], {})
            work = Work.objects.create(corpus=corpus, title_source=w['title'], title_target=m.get('译文标题', ''),
                                       author=m.get('作者', ''), sort_order=base_order + wi)
            position = 0
            for pi, para in enumerate(w['paragraphs']):
                paragraph = Paragraph.objects.create(work=work, seq=pi)
                seq_per_version = defaultdict(int)
                for gi, items in enumerate(para):
                    group = AlignGroup.objects.create(paragraph=paragraph, seq=gi, work=work, corpus=corpus,
                                                      position=position)
                    position += 1
                    for it in items:
                        v = versions[it['version']]
                        s = Segment.objects.create(group=group, version=v, seq=seq_per_version[v.id],
                                                   source=it['source'], target=it['target'], work=work,
                                                   corpus=corpus)
                        seq_per_version[v.id] += 1
                        new_seg_ids.append(s.id)
                        SegmentAnnotation.objects.bulk_create([
                            SegmentAnnotation(segment=s, value_id=values[(h, label)])
                            for h, labels in it['anns'].items() for label in labels])
        fts.index_segments(Segment.objects.filter(id__in=new_seg_ids).values_list('id', 'source', 'target'))
    return corpus, len(new_seg_ids)
