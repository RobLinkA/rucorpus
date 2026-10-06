import collections
import csv
import io

from django.db.models import Count
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from ninja import Query, Router, Schema
from ninja.errors import HttpError
from ninja.security import django_auth

from ..models import (AlignGroup, AnnotationValue, Corpus, Paragraph, SearchHistory, Segment, SegmentAnnotation,
                      Version, Work)
from ..search import SearchParams, all_group_ids, search
from ..ai.run import load_report
from ..serialize import ai_info, group_payloads
from ..text import Highlighter
from . import audit, require_editor

router = Router(tags=['public'])


def visible_corpora(request):
    qs = Corpus.objects.all()
    if not (request.user.is_authenticated and request.user.can_edit):
        qs = qs.filter(status=Corpus.Status.PUBLISHED)
    return qs


@router.get('/meta')
def meta(request):
    """Everything the UI needs to render labels, scopes and filters."""
    corpora = list(visible_corpora(request).prefetch_related('versions', 'works', 'annotation_groups__values'))
    ids = [c.id for c in corpora]
    groups_per_work = dict(AlignGroup.objects.filter(corpus_id__in=ids).values('work_id')
                           .annotate(n=Count('id')).values_list('work_id', 'n'))
    segs_per_version = dict(Segment.objects.filter(corpus_id__in=ids).values('version_id')
                            .annotate(n=Count('id')).values_list('version_id', 'n'))
    value_usage = dict(SegmentAnnotation.objects.filter(segment__corpus_id__in=ids).values('value_id')
                       .annotate(n=Count('id')).values_list('value_id', 'n'))
    ai_usage = dict(SegmentAnnotation.objects.filter(segment__corpus_id__in=ids, origin='ai').values('value_id')
                    .annotate(n=Count('id')).values_list('value_id', 'n'))
    return {'corpora': [{
        'id': c.id, 'name': c.name, 'description': c.description, 'source_lang': c.source_lang,
        'target_lang': c.target_lang, 'status': c.status,
        'versions': [{'id': v.id, 'kind': v.kind, 'lang': v.lang, 'label': v.label, 'person': v.person,
                      'bibliography': v.bibliography, 'segments': segs_per_version.get(v.id, 0)}
                     for v in c.versions.all()],
        'works': [{'id': w.id, 'title_source': w.title_source, 'title_target': w.title_target, 'author': w.author,
                   'groups': groups_per_work.get(w.id, 0)} for w in c.works.all()],
        'annotation_groups': [{'id': g.id, 'name': g.name, 'key': g.key, 'widget': g.widget,
                               'description': g.description,
                               'values': [{'id': v.id, 'label': v.label, 'usage': value_usage.get(v.id, 0),
                                           'ai_usage': ai_usage.get(v.id, 0),
                                           'defined_in_legacy': v.defined_in_legacy} for v in g.values.all()]}
                              for g in c.annotation_groups.all()],
    } for c in corpora]}


class SearchQuery(Schema):
    q: str = ''
    mode: str = 'morph'
    corpora: list[int] = []
    works: list[int] = []
    versions: list[int] = []
    values: list[int] = []
    exclude: list[int] = []
    logic: str = 'or'
    origin: str = 'all'
    preview: bool = False   # live count from the advanced-search page: not saved to history
    page: int = 1
    size: int = 20


def _params(request, sq: SearchQuery):
    allowed = set(visible_corpora(request).values_list('id', flat=True))
    corpora = [c for c in sq.corpora if c in allowed] or ([] if not sq.corpora else [-1])
    return SearchParams(q=sq.q.strip()[:300], mode='exact' if sq.mode == 'exact' else 'morph', corpora=corpora,
                        works=sq.works, versions=sq.versions, values=sq.values, exclude=sq.exclude,
                        logic=sq.logic if sq.logic in ('and', 'group') else 'or',
                        origin=sq.origin if sq.origin in ('human', 'ai') else 'all',
                        include_hidden=request.user.is_authenticated and request.user.can_edit)


def _summary(sq: SearchQuery):
    parts = []
    if sq.q:
        parts.append(f'“{sq.q}”' + ('（精确）' if sq.mode == 'exact' else ''))
    if sq.values:
        labels = list(AnnotationValue.objects.filter(id__in=sq.values).values_list('label', flat=True))
        parts.append({'and': '且', 'group': '/'}.get(sq.logic, '或').join(labels))
    if sq.exclude:
        parts.append('排除 ' + '、'.join(AnnotationValue.objects.filter(id__in=sq.exclude).values_list('label', flat=True)))
    if sq.origin in ('human', 'ai'):
        parts.append('仅人工标注' if sq.origin == 'human' else '仅 AI 标注')
    if sq.works:
        parts.append('、'.join(Work.objects.filter(id__in=sq.works).values_list('title_target', flat=True)[:4]))
    elif sq.corpora:
        parts.append('、'.join(Corpus.objects.filter(id__in=sq.corpora).values_list('name', flat=True)))
    if sq.versions:
        parts.append('、'.join(Version.objects.filter(id__in=sq.versions).values_list('label', flat=True)))
    return ' · '.join(parts)[:500]


@router.get('/search')
def do_search(request, sq: Query[SearchQuery]):
    p = _params(request, sq)
    size = min(max(sq.size, 5), 100)
    res = search(p, page=sq.page, page_size=size)
    hl = Highlighter(res['terms'], p.mode)
    history_id = None
    if request.user.is_authenticated and sq.page == 1 and not sq.preview:
        params = sq.dict(exclude={'page', 'size', 'preview'})
        last = SearchHistory.objects.filter(user=request.user).order_by('-created_at').first()
        if last and last.params == params:
            last.result_count = res['total_groups']
            last.save(update_fields=['result_count'])
            history_id = last.id
        else:
            history_id = SearchHistory.objects.create(user=request.user, params=params, summary=_summary(sq),
                                                      result_count=res['total_groups']).id
    return {
        'total': res['total_groups'], 'total_segments': res['total_segments'], 'page': sq.page, 'size': size,
        'facets': {k: [{'id': i, 'count': n} for i, n in v.items()] for k, v in res['facets'].items()},
        'results': group_payloads(res['group_ids'], request.user, res['hit_segments'], hl),
        'history_id': history_id,
    }


@router.get('/search/export')
def export_search(request, sq: Query[SearchQuery], format: str = 'xlsx'):
    p = _params(request, sq)
    group_ids, hits, _ = all_group_ids(p)
    return export_groups(group_ids, f'检索结果', format, hits=hits)


def export_groups(group_ids, title, fmt, hits=None, notes=None):
    """Export groups as a table: one row per segment (translation sentence pair)."""
    names = _label_maps()
    rows = [['语料库', '作品', '段落', '句组', '原文', '译本', '译文', '标注', '命中', '备注']]
    for chunk_start in range(0, len(group_ids), 500):
        for g in group_payloads(group_ids[chunk_start:chunk_start + 500], hit_segments=hits):
            for r in g['rows']:
                for s in r['segments']:
                    rows.append([names['corpus'][g['corpus_id']], names['work'][g['work_id']], g['paragraph_seq'] + 1,
                                 g['seq'] + 1, s['source'], names['version'][r['version_id']], s['target'],
                                 '；'.join(names['value'][v] for v in s['annotations']),
                                 '是' if s['hit'] else '', (notes or {}).get(g['id'], '')])
    if hits is None:
        rows = [r[:8] + r[9:] for r in rows]
    if not notes:
        rows = [r[:-1] for r in rows]
    return table_response(rows, title, fmt)


def _label_maps():
    return {
        'corpus': dict(Corpus.objects.values_list('id', 'name')),
        'work': {w.id: w.title_target or w.title_source for w in Work.objects.all()},
        'version': dict(Version.objects.values_list('id', 'label')),
        'value': {v.id: f'{v.group.name}:{v.label}' for v in AnnotationValue.objects.select_related('group')},
    }


def table_response(rows, title, fmt):
    from urllib.parse import quote
    if fmt == 'csv':
        buf = io.StringIO()
        buf.write('﻿')
        csv.writer(buf).writerows(rows)
        resp = HttpResponse(buf.getvalue().encode('utf-8'), content_type='text/csv; charset=utf-8')
        ext = 'csv'
    else:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        wb = Workbook()
        ws = wb.active
        ws.title = title[:30]
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True)
        widths = {'原文': 60, '译文': 50, '标注': 30, '备注': 30}
        for i, h in enumerate(rows[0], start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(h, 14)
            if h in widths:
                for cell in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                    for c in cell:
                        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.freeze_panes = 'A2'
        buf = io.BytesIO()
        wb.save(buf)
        resp = HttpResponse(buf.getvalue(),
                            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        ext = 'xlsx'
    resp['Content-Disposition'] = f"attachment; filename*=UTF-8''{quote(title)}.{ext}"
    return resp


@router.get('/read/{work_id}')
def read(request, work_id: int, page: int = 1, size: int = 40, around: int | None = None):
    work = get_object_or_404(Work, id=work_id, corpus__in=visible_corpora(request))
    size = min(max(size, 10), 200)
    total = AlignGroup.objects.filter(work=work).count()
    if around:
        pos = AlignGroup.objects.filter(id=around, work=work).values_list('position', flat=True).first()
        if pos is not None:
            page = pos // size + 1
    start = (max(page, 1) - 1) * size
    ids = list(AlignGroup.objects.filter(work=work, position__gte=start, position__lt=start + size)
               .order_by('position').values_list('id', flat=True))
    return {'work_id': work.id, 'corpus_id': work.corpus_id, 'total': total, 'page': page, 'size': size,
            'paragraphs': Paragraph.objects.filter(work=work).count(),
            'groups': group_payloads(ids, request.user)}


@router.get('/groups/{group_id}')
def group_detail(request, group_id: int):
    g = get_object_or_404(AlignGroup.objects.select_related('paragraph'), id=group_id,
                          corpus__in=visible_corpora(request))
    payload = group_payloads([g.id], request.user)[0]
    # context: the whole paragraph when short, else a window around the group
    para_groups = AlignGroup.objects.filter(paragraph_id=g.paragraph_id).order_by('seq')
    if para_groups.count() > 25:
        para_groups = AlignGroup.objects.filter(work_id=g.work_id, position__gte=g.position - 8,
                                                position__lte=g.position + 8).order_by('position')
    context = group_payloads(list(para_groups.values_list('id', flat=True)))
    nav = dict(AlignGroup.objects.filter(work_id=g.work_id, position__in=[g.position - 1, g.position + 1])
               .values_list('position', 'id'))
    payload.update({
        'context': [{'id': c['id'], 'source': c['source'], 'is_current': c['id'] == g.id, 'paragraph_seq': c['paragraph_seq'],
                     'targets': {r['version_id']: ' '.join(s['target'] for s in r['segments']) for r in c['rows']}}
                    for c in context],
        'prev': nav.get(g.position - 1), 'next': nav.get(g.position + 1),
        'work_total': AlignGroup.objects.filter(work_id=g.work_id).count(),
    })
    return payload


class AnnotationsIn(Schema):
    values: list[int]
    confirm_ai: bool = False


@router.put('/segments/{segment_id}/annotations', auth=django_auth)
def set_annotations(request, segment_id: int, data: AnnotationsIn):
    require_editor(request)
    seg = get_object_or_404(Segment, id=segment_id)
    valid = AnnotationValue.objects.filter(id__in=data.values, group__corpus_id=seg.corpus_id) \
        .select_related('group')
    if len(valid) != len(set(data.values)):
        raise HttpError(400, '标注项不属于该语料库')
    per_group = collections.Counter(v.group_id for v in valid if v.group.single_choice)
    if any(n > 1 for n in per_group.values()):
        raise HttpError(400, '单选标注组只能选一个值')
    before = set(SegmentAnnotation.objects.filter(segment=seg).values_list('value_id', flat=True))
    after = {v.id for v in valid}
    SegmentAnnotation.objects.filter(segment=seg, value_id__in=before - after).delete()
    SegmentAnnotation.objects.bulk_create([SegmentAnnotation(segment=seg, value_id=v) for v in after - before])
    if data.confirm_ai:  # an editor confirming AI annotations turns them into human ones
        SegmentAnnotation.objects.filter(segment=seg, origin='ai', value_id__in=after).update(origin='human')
    if before != after:
        labels = {v.id: f'{v.group.name}:{v.label}' for v in AnnotationValue.objects.filter(id__in=before | after)
                  .select_related('group')}
        audit(request, 'annotate', f'segment:{seg.id}',
              f"句对 {seg.id}：+{'、'.join(labels[i] for i in after - before) or '无'} "
              f"−{'、'.join(labels[i] for i in before - after) or '无'}",
              added=sorted(after - before), removed=sorted(before - after))
    return {'id': seg.id, 'annotations': sorted(after), 'ai': ai_info([seg.id]).get(seg.id, {})}


@router.get('/ai/report')
def ai_report(request):
    return load_report() or {}


@router.get('/stats')
def stats(request):
    vis = visible_corpora(request)
    return {
        'ai_annotations': SegmentAnnotation.objects.filter(segment__corpus__in=vis, origin='ai').count(),
        'corpora': vis.count(),
        'works': Work.objects.filter(corpus__in=vis).count(),
        'groups': AlignGroup.objects.filter(corpus__in=vis).count(),
        'segments': Segment.objects.filter(corpus__in=vis).count(),
        'annotations': SegmentAnnotation.objects.filter(segment__corpus__in=vis).count(),
    }
